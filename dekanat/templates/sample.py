"""Учебные образцы шаблонов «Заявление о переводе» и «Заявление о зачислении» — для знакомства с форматом и для тестов. Это не официальные формы вуза.

Часть плейсхолдеров намеренно разбита на несколько runs с разным форматированием — так Word сохраняет текст после правок.
Администратор загружает свой шаблон (с шапкой и оформлением вуза) тем же способом: в тексте достаточно записать `{{ student_name }}` и т. п.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def build_sample_template(path: "str | Path") -> Path:
    path = Path(path)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)

    section = doc.sections[0]
    section.header.paragraphs[0].text = "Образец. Заявление о переводе (учебная форма, не официальный бланк)"
    section.header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    section.footer.paragraphs[0].text = "Дата подачи: {{ application_date }}"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("Декану факультета\n")
    p.add_run("от студента ")
    r = p.add_run("{{ stu")  # плейсхолдер разбит на два run с разным начертанием
    r.bold = True
    r2 = p.add_run("dent_name }}")
    r2.bold = True
    r2.italic = True

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = title.add_run("ЗАЯВЛЕНИЕ")
    t.bold = True

    doc.add_paragraph(
        "Прошу рассмотреть возможность моего перевода из {{ current_university }}, где я обучаюсь на {{ current_course }} курсе "
        "по направлению подготовки «{{ current_specialty }}» ({{ education_form }} форма обучения), "
        "на направление подготовки «{{ target_specialty }}»."
    )
    doc.add_paragraph("Академические задолженности: {{ debts }}.")

    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    rows = [("Паспортные данные", "{{ passport_data }}"), ("Дата заявления", "{{ application_date }}"), ("Подпись студента", "________________")]
    for row, (a, b) in zip(table.rows, rows):
        row.cells[0].text = a
        row.cells[1].text = b

    doc.add_paragraph()
    doc.add_paragraph("К заявлению прилагается справка об обучении.")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def build_enrollment_template(path: "str | Path") -> Path:
    """Образец «Заявление о зачислении»: те же метки, что собирает анкета (сценарий «зачисление»), оформление — как у образца о переводе."""
    path = Path(path)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)

    section = doc.sections[0]
    section.header.paragraphs[0].text = "Образец. Заявление о зачислении (учебная форма, не официальный бланк)"
    section.header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    section.footer.paragraphs[0].text = "Дата подачи: {{ application_date }}"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("Декану факультета\n")
    p.add_run("от ")
    p.add_run("{{ student_name }}").bold = True

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("ЗАЯВЛЕНИЕ").bold = True

    doc.add_paragraph(
        "Прошу зачислить меня на направление подготовки «{{ target_specialty }}» ({{ education_form }} форма обучения)."
    )
    doc.add_paragraph(
        "Сведения о прежнем обучении: {{ current_university }}, {{ current_course }} курс, направление подготовки «{{ current_specialty }}»."
    )
    doc.add_paragraph("Академические задолженности: {{ debts }}.")

    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    rows = [("Паспортные данные", "{{ passport_data }}"), ("Дата заявления", "{{ application_date }}"), ("Подпись заявителя", "________________")]
    for row, (a, b) in zip(table.rows, rows):
        row.cells[0].text = a
        row.cells[1].text = b

    doc.add_paragraph()
    doc.add_paragraph("К заявлению прилагаются документы согласно перечню деканата.")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
