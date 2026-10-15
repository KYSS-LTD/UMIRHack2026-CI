"""Разбор Положения на нумерованные пункты («3.4», «п. 3.4», «Статья 5»...).

Поддерживаются .md / .txt / .docx / .pdf. Нумерация пунктов определяется автоматически:
вложенные пункты («3.2», «3.2.1»), заголовки разделов («3. Справка об обучении» и «## 3. ...»),
подпункты («а)», «1)», «-») присоединяются к своему пункту.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Optional

from ..models import Clause

_MD_HEAD = re.compile(r"^\s{0,3}#{1,6}\s*(.+?)\s*#*\s*$")
_CLAUSE = re.compile(r"^\s*(?:(?:п(?:ункт)?|ст(?:атья)?|§)\.?\s*)?(\d{1,3}(?:\.\d{1,3})*)\s*([.)])?\s+(\S.*)$", re.IGNORECASE)
_CLAUSE_PREFIXED = re.compile(r"^\s*(?:п(?:ункт)?|ст(?:атья)?|§)\.?\s*(\d{1,3}(?:\.\d{1,3})*)\s*[.):]?\s*(.*)$", re.IGNORECASE)
_SUBITEM = re.compile(r"^\s*(?:[а-яa-z]\)|\d{1,2}\)|[-–—•*·]\s)\s*(\S.*)$")
_SECTION_WORD = re.compile(r"^\s*(?:раздел|глава)\s+(\d{1,3}|[IVX]+)[.:]?\s*(.*)$", re.IGNORECASE)
_EMPH = re.compile(r"(\*\*|__|\*|_)(.+?)\1")


def _strip_markdown(line: str) -> str:
    line = _EMPH.sub(lambda m: m.group(2), line)
    return line.replace(" ", " ").rstrip()


def _is_heading_text(text: str) -> bool:
    """Короткая строка без завершающей точки — вероятнее заголовок, чем пункт."""
    t = text.strip()
    return len(t) <= 90 and not t.endswith((".", ";", ":", "!", "?")) and len(t.split()) <= 12


def parse_regulation_text(text: str) -> list[Clause]:
    lines = [_strip_markdown(l) for l in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    all_numbers = {m.group(1) for m in (_CLAUSE.match(l.strip()) for l in lines) if m and ("." in m.group(1) or m.group(2))}
    headings: dict[str, str] = {}  # «3» -> «Справка об обучении»
    md_stack: list[str] = []
    clauses: list[Clause] = []
    cur: Optional[dict] = None

    def flush():
        nonlocal cur
        if cur and cur["text"].strip():
            clauses.append(Clause(cur["number"], re.sub(r"\s+", " ", cur["text"]).strip(), list(cur["section"]), True))
        cur = None

    def section_for(number: str) -> list[str]:
        parts = number.split(".")
        out = []
        for i in range(1, len(parts)):
            h = headings.get(".".join(parts[:i]))
            if h:
                out.append(h)
        if not out and md_stack:
            out = list(md_stack)
        return out

    for raw in lines:
        line = raw.strip()
        if not line:
            if cur:
                cur["blank"] = True
            continue
        md = _MD_HEAD.match(line)
        if md:
            flush()
            title = md.group(1).strip()
            m = _CLAUSE.match(title)
            if m and (m.group(2) or "." in m.group(1)):
                headings[m.group(1)] = m.group(3).strip()
                md_stack = [m.group(3).strip()]
            else:
                md_stack = [title]
            continue
        sec = _SECTION_WORD.match(line)
        if sec:
            flush()
            headings[sec.group(1)] = (sec.group(2) or line).strip()
            continue
        m = _CLAUSE.match(line)
        if m:
            number, dot, rest = m.group(1), m.group(2), m.group(3).strip()
            depth = number.count(".") + 1
            prefixed = bool(re.match(r"^\s*(?:п|пункт|ст|статья|§)", line, re.IGNORECASE))
            valid = depth > 1 or bool(dot) or prefixed
            if valid and depth == 1 and not dot and not prefixed:
                valid = False
            # «2024 год»: слишком большое число без разделителя — не пункт
            if valid and depth == 1 and int(number) > 200:
                valid = False
            if valid:
                flush()
                if depth == 1 and _is_heading_text(rest) and any(n.startswith(number + ".") for n in all_numbers):
                    headings[number] = rest
                    continue
                cur = {"number": number, "text": rest, "section": section_for(number), "blank": False}
                continue
        sub = _SUBITEM.match(line)
        if cur is not None and (sub or not cur.get("blank")):
            cur["text"] += " " + (line if not sub else line)
            continue
        if cur is not None and cur.get("blank"):
            # абзац после пустой строки — продолжение пункта, если не начинается с новой нумерации
            cur["text"] += " " + line
            cur["blank"] = False
            continue
        # текст до первого пункта (преамбула) — в базу знаний не включаем
    flush()
    if not clauses:
        return _paragraph_fallback(text)
    return clauses


def _paragraph_fallback(text: str) -> list[Clause]:
    """Положение без нумерации: делим на абзацы и нумеруем сами («§1», «§2»...)."""
    paras = [re.sub(r"\s+", " ", p).strip() for p in re.split(r"\n\s*\n", text) if len(p.strip()) > 20]
    return [Clause(f"§{i}", p, [], False) for i, p in enumerate(paras, 1)]


def _read_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("Для .docx нужен пакет python-docx (pip install python-docx)") from exc
    document = docx.Document(str(path))
    counters: dict[str, list[int]] = {}
    out = []
    for p in document.paragraphs:
        text = p.text.strip()
        if not text:
            out.append("")
            continue
        num_pr = p._p.pPr.numPr if p._p.pPr is not None else None
        if num_pr is not None and num_pr.numId is not None and not re.match(r"^\s*(?:п\.?\s*)?\d+(?:\.\d+)*[.)]?\s", text):
            # автонумерация Word: восстанавливаем «1.2.» по уровням списка
            nid = str(num_pr.numId.val)
            lvl = int(num_pr.ilvl.val) if num_pr.ilvl is not None else 0
            c = counters.setdefault(nid, [])
            while len(c) <= lvl:
                c.append(0)
            c[lvl] += 1
            del c[lvl + 1:]
            text = ".".join(str(x) for x in c[: lvl + 1]) + ". " + text
        if (p.style.name or "").lower().startswith("heading"):
            text = "## " + text
        out.append(text)
    for table in document.tables:
        for row in table.rows:
            out.append(" ".join(c.text.strip() for c in row.cells))
    return "\n".join(out)


def _read_pdf(path: Path) -> str:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    return "\n".join(pdf[i].get_textpage().get_text_range() for i in range(len(pdf)))


def read_regulation_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _read_docx(path)
    if suffix == ".pdf":
        return _read_pdf(path)
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def load_regulation(path: Path) -> list[Clause]:
    if not path.exists():
        raise FileNotFoundError(f"Файл Положения не найден: {path}")
    return parse_regulation_text(read_regulation_file(path))
