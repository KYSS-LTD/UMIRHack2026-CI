"""База знаний и шаблоны «как у пользователя»: сотрудник загружает документ, секретарь отвечает по нему; шаблон — анкета — готовый документ."""

import io
import re
import zipfile

import pytest
from docx import Document
from fastapi.testclient import TestClient

from conftest import admin_login
from dekanat.dialog.intents import detect_intent
from dekanat.templates.sample import build_sample_template
from dekanat.web.app import create_app

HOSTEL = """# Правила проживания в общежитии (учебный документ)

## 1. Заселение

1.1. Студент подаёт заявление на заселение в общежитие не позднее 5 рабочих дней до начала семестра.

1.2. Для заселения необходимо предъявить паспорт, студенческий билет и медицинскую справку формы 086.

## 2. Проживание

2.1. Плата за проживание в общежитии составляет 1500 рублей в месяц и вносится до 10 числа.

2.2. Посторонние лица допускаются в общежитие только до 22:00.
"""


@pytest.fixture()
def app(tmp_settings):
    return create_app(tmp_settings)


@pytest.fixture()
def admin(app):
    c = TestClient(app)
    admin_login(c, app)
    return c


def upload_kb(admin, text=HOSTEL, name="Правила общежития", filename="hostel.md"):
    r = admin.post("/api/admin/kb", files={"file": (filename, text.encode("utf-8"), "text/markdown")}, data={"name": name})
    assert r.status_code == 200, r.text
    return r.json()


def answer(app, question):
    """Вопрос без документа на столе — так его получает сайт (сообщение вне анкеты -> проверка)."""
    return app.state.svc.analyze(None, "", question).answer


class TestUploadedDocumentAnswers:
    def test_new_document_is_used_at_once_with_its_own_citation(self, app, admin):
        assert answer(app, "Сколько стоит проживание в общежитии?").verdict == "no_answer"
        upload_kb(admin)
        a = answer(app, "Сколько стоит проживание в общежитии?")
        assert a.verdict == "ok" and "1500" in a.text and a.sources[0].ref.endswith(":2.1")
        assert "Правила общежития" in a.text  # ссылка называет документ: пункты разных документов не путаются

    @pytest.mark.parametrize("question, fragment", [
        ("До какого времени допускаются посторонние в общежитие?", "22:00"),
        ("За сколько дней подать заявление на заселение?", "5 рабочих дней"),
        ("Сколько стоит проживание в общежитии?", "1500"),
    ])
    def test_questions_find_the_clause(self, app, admin, question, fragment):
        upload_kb(admin)
        a = answer(app, question)
        assert a.verdict == "ok" and fragment in a.text, a.text

    def test_a_cost_clause_about_another_subject_does_not_answer_tuition(self, app, admin):
        """«Сколько стоит обучение?» — пример из ТЗ с ответом «нет данных»: плата за общежитие им не является."""
        upload_kb(admin)
        a = answer(app, "Сколько стоит обучение?")
        assert a.verdict == "no_answer" and not a.sources
        assert a.text == "В документе и в Положении нет информации о стоимости обучения. Обратитесь в деканат."

    def test_disable_enable_replace_delete(self, app, admin):
        doc = upload_kb(admin)["id"]
        assert admin.patch(f"/api/admin/kb/{doc}", json={"active": False}).status_code == 200
        assert answer(app, "Сколько стоит проживание в общежитии?").verdict == "no_answer"
        assert admin.patch(f"/api/admin/kb/{doc}", json={"active": True}).status_code == 200
        assert "1500" in answer(app, "Сколько стоит проживание в общежитии?").text
        r = admin.put(f"/api/admin/kb/{doc}", files={"file": ("h2.md", HOSTEL.replace("1500", "2000").encode("utf-8"), "text/markdown")})
        assert r.status_code == 200 and r.json()["version"] == 2
        text = answer(app, "Сколько стоит проживание в общежитии?").text
        assert "2000" in text and "1500" not in text
        assert [v["version"] for v in admin.get(f"/api/admin/kb/{doc}/history").json()["versions"]] == [2, 1]
        assert admin.delete(f"/api/admin/kb/{doc}").status_code == 200
        assert answer(app, "Сколько стоит проживание в общежитии?").verdict == "no_answer"
        assert answer(app, "Сколько дней рассматривается заявление?").sources[0].ref == "6.1"  # основное Положение на месте

    def test_docx_document_is_read(self, app, admin):
        buf = io.BytesIO()
        d = Document()
        d.add_paragraph("Положение о стипендии")
        d.add_paragraph("1.1. Академическая стипендия назначается студентам, сдавшим сессию без троек.")
        d.add_paragraph("1.2. Размер академической стипендии составляет 3000 рублей в месяц.")
        d.save(buf)
        r = admin.post("/api/admin/kb", files={"file": ("stip.docx", buf.getvalue(), "application/octet-stream")}, data={"name": "Положение о стипендии"})
        assert r.status_code == 200 and r.json()["clauses"] == 2
        assert "3000" in answer(app, "Какой размер стипендии?").text


class TestUploadChecks:
    def test_files_without_numbering_are_accepted_with_a_notice(self, admin):
        r = admin.post("/api/admin/kb", files={"file": ("lib.txt", "Заниматься можно до 21:00.\n\nКниги выдаются на 14 дней.".encode(), "text/plain")}, data={"name": "Библиотека"})
        assert r.status_code == 200 and r.json()["clauses"] == 2
        assert "нумерации пунктов" in r.json()["notice"] and "§1" in r.json()["notice"]
        assert "notice" not in upload_kb(admin)  # с нумерацией предупреждения нет

    def test_replace_with_unnumbered_file_also_warns(self, admin):
        doc = upload_kb(admin)["id"]
        r = admin.put(f"/api/admin/kb/{doc}", files={"file": ("x.txt", "Заниматься в библиотеке можно до девяти вечера. Книги выдаются на две недели.".encode("utf-8"), "text/plain")})
        assert r.status_code == 200 and "нумерации" in r.json()["notice"]

    @pytest.mark.parametrize("name, data, status", [("пустой.txt", b"", 400), ("картинка.png", b"\x89PNG\r\n\x1a\n" + b"0" * 40, 400)])
    def test_bad_files_are_refused(self, admin, name, data, status):
        r = admin.post("/api/admin/kb", files={"file": (name, data, "application/octet-stream")}, data={"name": name})
        assert r.status_code == status

    def test_same_name_is_refused(self, admin):
        upload_kb(admin)
        r = admin.post("/api/admin/kb", files={"file": ("h.md", HOSTEL.encode(), "text/markdown")}, data={"name": "Правила общежития"})
        assert r.status_code == 400 and "уже есть" in r.json()["detail"]


class TestQuestionsAboutSubmittingAreNotTheQuestionnaire:
    @pytest.mark.parametrize("text", ["За сколько дней подать заявление на заселение?", "Когда подавать заявление?", "Где подать заявление на перевод?", "Сколько дней рассматривается заявление о переводе?"])
    def test_rule_questions_do_not_offer_the_form(self, text):
        assert detect_intent(text) is None

    @pytest.mark.parametrize("text", ["Я хочу перевестись", "Хочу подать заявление на перевод", "Мне нужно оформить заявление", "Как к вам поступить?"])
    def test_wishes_still_do(self, text):
        assert detect_intent(text) is not None


# ============================================================================ шаблоны
def docx_text(data: bytes) -> str:
    z = zipfile.ZipFile(io.BytesIO(data))
    return "\n".join(" ".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", z.read(n).decode("utf-8"))) for n in z.namelist() if n.startswith("word/") and n.endswith(".xml"))


def template_with_custom_key() -> bytes:
    d = Document()
    d.sections[0].header.paragraphs[0].text = "От {{ student_name }}"
    p = d.add_paragraph()
    p.add_run("Прошу перевести из {{ current_")  # метка разорвана, как после правок в Word
    p.add_run("university }} на «{{ target_specialty }}».")
    d.add_paragraph("Номер зачётки: {{ номер_зачётки }}; ещё: {{ record_book }}")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def test_template_with_split_and_custom_keys_through_the_whole_questionnaire(app, admin, tmp_path):
    r = admin.post("/api/admin/templates", files={"file": ("t.docx", template_with_custom_key(), "application/octet-stream")}, data={"name": "Проба", "template_id": "", "replace": "false"})
    assert r.status_code == 200
    keys = {p["key"] for p in r.json()["placeholders"]}
    assert {"student_name", "current_university", "target_specialty", "номер_зачётки", "record_book"} <= keys
    tpl = r.json()["template_id"]

    c = TestClient(app)
    c.get("/api/assistant/state")
    say = lambda t: c.post("/api/assistant/message", json={"text": t, "defer": True}).json()
    m = say("Я хочу перевестись")
    m = say("да")
    own = {"student_name": "Иванов Иван Иванович", "current_university": "Южный федеральный университет", "current_course": "2", "current_specialty": "09.03.01 Информатика",
           "education_form": "очная", "has_debts": "нет", "request_type": "перевод", "target_specialty": "Прикладная информатика", "passport_data": "пропустить"}
    for _ in range(14):
        if m["reply"]["kind"] == "review":
            break
        m = say(own[m["reply"]["question"]])
    say("верно")
    g = c.post("/api/assistant/generate", json={"template_id": tpl}).json()
    text = docx_text(c.get(g["download"]["url"]).content)
    assert "Иванов Иван Иванович" in text and "Южный федеральный университет" in text and "Прикладная информатика" in text
    assert "{{" not in text and "________" in text
    # студент видит читаемые названия, а не имена переменных
    assert set(g["unfilled"]) == {"номер зачётки", "record book"}
    assert "record_book" not in g["reply"]["text"] and "номер зачётки" in g["reply"]["text"]
