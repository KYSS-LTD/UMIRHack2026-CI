"""Устойчивость: повреждённое состояние, срок хранения паспортных данных, мусорный ввод, ограничение числа сессий, модель на уровне приложения."""

import json
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import pytest
from fastapi.testclient import TestClient

from conftest import admin_login
from dekanat.dialog import IDLE, REVIEW, SCREENING, DialogContext, DialogMachine
from dekanat.templates.sample import build_sample_template
from dekanat.web.app import create_app

ANSWERS = ["Иванов Иван Иванович", "Южный федеральный университет", "2", "09.03.01 Информатика", "очная", "нет"]


def say(c, *texts):
    out = None
    for t in texts:
        out = c.post("/api/assistant/message", json={"text": t})
        assert out.status_code == 200, (t, out.text)
    return out.json()


# ============================================================================ повреждённое состояние
@pytest.mark.parametrize("raw", [
    {"state": "screening", "current": None, "data": {"student_name": "Иванов Иван"}},
    {"state": "screening", "current": "removed_question", "scenario": "transfer"},
    {"state": "review", "data": {}},
    {"state": "document_generation", "data": []},
    {"state": "screening", "data": "oops", "skipped": "oops", "attempts": [], "scenario": "weird"},
    {"state": 5, "data": None},
])
def test_corrupted_state_never_crashes_the_dialog(raw):
    ctx = DialogContext.from_dict(raw)
    m = DialogMachine()
    for text in ("Иванов Иван Иванович", "продолжить", "назад", "исправить", "верно", "Я хочу перевестись"):
        reply = m.handle(ctx, text)
        assert reply.state in ("idle", "screening", "review", "document_generation", "completed")
        json.dumps(ctx.to_dict())  # состояние остаётся сериализуемым


def test_screening_without_current_question_continues_from_the_first_unanswered():
    ctx = DialogContext.from_dict({"state": "screening", "scenario": "transfer", "data": {"request_type": "transfer", "student_name": "Иванов Иван"}})
    reply = DialogMachine().handle(ctx, "продолжить")
    assert ctx.state == SCREENING and ctx.current == "current_university" and reply.question == "current_university"


def test_review_with_no_data_falls_back_to_idle():
    ctx = DialogContext.from_dict({"state": "review"})
    DialogMachine().handle(ctx, "верно")
    assert ctx.state == IDLE


# ============================================================================ паспортные данные: срок хранения
def test_passport_is_dropped_after_its_retention_time():
    m = DialogMachine()
    ctx = DialogContext()
    for t in ("Я хочу перевестись", "да", *ANSWERS, "Прикладная информатика", "1234 567890"):
        m.handle(ctx, t)
    assert ctx.state == REVIEW and ctx.data["passport_data"] == "серия 1234 № 567890" and ctx.passport_at > 0
    assert DialogMachine.expire_sensitive(ctx, 3600) is False  # срок ещё не вышел
    assert DialogMachine.expire_sensitive(ctx, 3600, now=ctx.passport_at + 3601) is True
    assert "passport_data" not in ctx.data and "passport_data" in ctx.skipped and ctx.passport_at == 0
    assert DialogMachine.template_data(ctx)["passport_data"] is None  # в документ пойдёт «не заполнено», а не старое значение


def test_passport_expiry_through_the_api(tmp_settings, tmp_path):
    app = create_app(tmp_settings)
    admin = TestClient(app)
    admin_login(admin, app)
    admin.post("/api/admin/templates", files={"file": ("p.docx", build_sample_template(tmp_path / "t.docx").read_bytes(), "application/octet-stream")}, data={"template_id": "perevod"})
    c = TestClient(app)
    say(c, "Я хочу перевестись", "да", *ANSWERS, "Прикладная информатика", "1234 567890")
    sid = c.cookies.get("dekanat_sid")
    state = app.state.storage.sessions.load_state(sid)
    assert state["data"]["passport_data"]
    state["passport_at"] = time.time() - 7200  # введены два часа назад (срок — 60 минут)
    app.state.storage.sessions.save_state(sid, state)
    say(c, "верно")
    j = c.post("/api/assistant/generate", json={"template_id": "perevod"}).json()
    assert "паспортные данные" in j["unfilled"]
    assert "passport_data" not in app.state.storage.sessions.load_state(sid)["data"]


# ============================================================================ мусорный ввод
JUNK = ["\x00\x01\x02", "а" * 1999, "😀😀😀", "'; DROP TABLE tickets; --", "<script>alert(1)</script>", "{{ student_name }}", "../../etc/passwd", "\n\n\n", "   ", "?" * 50, "ё" * 300]


@pytest.mark.parametrize("text", JUNK)
def test_junk_messages_never_cause_a_server_error(tmp_settings, text):
    c = TestClient(create_app(tmp_settings))
    for prefix in ((), ("Я хочу перевестись", "да")):
        r = say(c, *prefix, text)
        assert r["reply"]["state"] in ("idle", "screening")
    assert c.get("/api/assistant/state").status_code == 200


def test_junk_filenames_and_files(tmp_settings):
    c = TestClient(create_app(tmp_settings))
    c.get("/api/assistant/state")
    for name in ("../../x.png", "a\nb.png", "x" * 400 + ".png", "con.png", ".png", "я.pdf"):
        r = c.post("/api/assistant/document", files={"file": (name, b"not an image at all", "image/png")})
        assert r.status_code == 400, (name, r.status_code)
    assert not list(tmp_settings.root.rglob("x.png"))


def test_templates_without_placeholders_are_accepted_but_reported(tmp_settings, tmp_path):
    from docx import Document

    app = create_app(tmp_settings)
    admin = TestClient(app)
    admin_login(admin, app)
    doc = Document()
    doc.add_paragraph("Просто текст без переменных")
    path = tmp_path / "plain.docx"
    doc.save(str(path))
    r = admin.post("/api/admin/templates", files={"file": ("plain.docx", path.read_bytes(), "application/octet-stream")})
    assert r.status_code == 200 and r.json()["placeholders"] == []


# ============================================================================ ограничение числа сессий
def test_mass_session_creation_is_limited(tmp_settings):
    tmp_settings.raw["server"]["max_new_sessions_per_minute"] = 3
    app = create_app(tmp_settings)
    codes = [TestClient(app).get("/api/assistant/state").status_code for _ in range(5)]
    assert codes == [200, 200, 200, 429, 429]


def test_zero_means_no_limit_on_new_sessions(tmp_settings):
    tmp_settings.raw["server"]["max_new_sessions_per_minute"] = 0
    app = create_app(tmp_settings)
    assert [TestClient(app).get("/api/assistant/state").status_code for _ in range(5)] == [200] * 5


def test_existing_session_still_works_when_the_limit_is_reached(tmp_settings):
    tmp_settings.raw["server"]["max_new_sessions_per_minute"] = 1
    app = create_app(tmp_settings)
    c = TestClient(app)
    assert c.get("/api/assistant/state").status_code == 200
    assert TestClient(app).get("/api/assistant/state").status_code == 429
    assert say(c, "привет")["session"] == "ok"  # живая сессия новых создавать не требует


# ============================================================================ основное Положение отключено
def test_disabling_the_base_regulation_does_not_break_answers(tmp_settings):
    app = create_app(tmp_settings)
    admin = TestClient(app)
    admin_login(admin, app)
    assert admin.patch("/api/admin/kb/base", json={"active": False}).status_code == 200
    assert len(admin.get("/api/regulation").json()["clauses"]) == 0
    r = admin.post("/api/analyze", data={"question": "Сколько дней рассматривается заявление?"})
    assert r.status_code == 200 and r.json()["answer"]["verdict"] == "no_answer" and r.json()["answer"]["sources"] == []
    assert admin.post("/api/analyze", data={"question": "Хватит ли справки?"}).status_code == 200
    admin.patch("/api/admin/kb/base", json={"active": True})
    assert len(admin.get("/api/regulation").json()["clauses"]) == 25


# ============================================================================ локальная модель на уровне всего приложения
def llm_settings(settings, url):
    settings.raw["USE_LOCAL_LLM"] = True
    settings.raw["llm"].update(ollama_url=url, timeout=3, cooldown_seconds=60)
    return settings


def test_app_works_when_the_enabled_model_is_unreachable(tmp_settings):
    app = create_app(llm_settings(tmp_settings, "http://127.0.0.1:9"))
    c = TestClient(app)
    h = c.get("/api/health").json()
    assert h["llm"]["enabled"] is True and h["llm"]["available"] is False
    assert c.post("/api/analyze", data={"question": "Как долго готовится справка об обучении?"}).json()["answer"]["verdict"] == "no_answer"
    assert c.post("/api/analyze", data={"question": "Сколько дней рассматривается заявление?"}).json()["answer"]["sources"][0]["ref"] == "6.1"
    say(c, "помогите, хочу к вам из другого университета")  # модель недоступна — это обычное сообщение
    r = say(c, "Я хочу перевестись", "да", *ANSWERS[:5], "Да не, долгов вроде нет, всё закрыл")
    assert r["reply"]["kind"] == "clarify"  # без модели просто уточняем
    assert say(c, "нет")["reply"]["question"] == "target_specialty"
    assert app.state.svc.llm.stats.ok == 0


class _FakeOllama(BaseHTTPRequestHandler):
    prompts: list = []

    def log_message(self, *a):
        pass

    def _send(self, body):
        raw = json.dumps(body).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(raw)))
        self.end_headers()
        self.wfile.write(raw)

    def do_GET(self):
        self._send({"models": [{"name": "qwen2.5:3b-instruct"}]})

    def do_POST(self):
        body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        user = body["messages"][1]["content"]
        _FakeOllama.prompts.append(user)
        if "превратить свободный ответ" in user:
            content = {"status": "clear", "value": False, "details": None, "clarification": None}
        elif "определить намерение" in user:
            content = {"intent": "transfer", "confidence": "high"}
        else:
            content = {"answerable": False, "clause": None}
        self._send({"message": {"content": json.dumps(content)}})


def test_app_uses_a_working_model_only_for_allowed_tasks(tmp_settings):
    _FakeOllama.prompts = []
    server = ThreadingHTTPServer(("127.0.0.1", 0), _FakeOllama)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    try:
        app = create_app(llm_settings(tmp_settings, f"http://127.0.0.1:{server.server_address[1]}"))
        c = TestClient(app)
        assert c.get("/api/health").json()["llm"]["available"] is True
        r = say(c, "помогите, хочу к вам из другого университета")
        assert r["reply"]["kind"] == "offer"  # модель распознала намерение, но анкета ждёт согласия
        say(c, "да", *ANSWERS[:5])
        r = say(c, "Да не, долгов вроде нет, всё закрыл")
        assert r["reply"]["question"] == "target_specialty"  # однозначный ответ модели принят после проверки
        st = app.state.storage.sessions.load_state(c.cookies.get("dekanat_sid"))["data"]
        assert st["has_debts"] is False
        assert set(app.state.svc.llm.task_counts) <= {"detect_intent", "normalize_answer", "judge_clauses", "extract_fields", "explain_answer"}
        assert app.state.svc.llm.task_counts == {"detect_intent": 1, "normalize_answer": 1}
        assert all("<данные>" in p for p in _FakeOllama.prompts)  # пользовательский текст всегда внутри блока данных
    finally:
        server.shutdown()


# ============================================================================ регрессии, найденные при ревью
def _at_review():
    m = DialogMachine()
    ctx = DialogContext()
    for t in ("Я хочу перевестись", "да", *ANSWERS, "Прикладная информатика", "1234 567890"):
        m.handle(ctx, t)
    assert ctx.state == REVIEW
    return m, ctx


def test_skipping_an_optional_field_during_edit_returns_to_review():
    m, ctx = _at_review()
    m.handle(ctx, "исправить паспорт")
    assert ctx.editing and ctx.current == "passport_data"
    r = m.handle(ctx, "пропустить")
    assert ctx.state == REVIEW and not ctx.editing and ctx.resume is None and r.kind == "review"
    assert "passport_data" in ctx.skipped and "passport_data" not in ctx.data and ctx.passport_at == 0
    assert ctx.skipped.count("passport_data") == 1
    m.handle(ctx, "исправить курс")  # дальнейшие правки работают как обычно
    m.handle(ctx, "3")
    assert ctx.state == REVIEW and ctx.data["current_course"] == "3"


def test_editing_in_the_middle_of_screening_resumes_the_current_question():
    m = DialogMachine()
    ctx = DialogContext()
    for t in ("Я хочу перевестись", "да", *ANSWERS):
        m.handle(ctx, t)
    assert ctx.current == "target_specialty"
    m.handle(ctx, "исправить курс")
    assert ctx.editing and ctx.current == "current_course" and ctx.resume == "target_specialty"
    m.handle(ctx, "4")
    assert ctx.current == "target_specialty" and not ctx.editing and ctx.resume is None and ctx.data["current_course"] == "4"
    m.handle(ctx, "исправить паспорт")  # паспорта ещё нет в ответах — править пока нечего
    assert ctx.state == SCREENING and ctx.current == "target_specialty"


def test_other_commands_cancel_the_choice_of_a_field_to_fix():
    m, ctx = _at_review()
    m.handle(ctx, "исправить")
    assert ctx.pick_fix
    m.handle(ctx, "отмена")
    assert not ctx.pick_fix
    m2, ctx2 = DialogMachine(), DialogContext()
    for t in ("Я хочу перевестись", "да", "Иванов Иван Иванович"):
        m2.handle(ctx2, t)
    m2.handle(ctx2, "исправить")
    assert ctx2.pick_fix
    m2.handle(ctx2, "назад")
    assert not ctx2.pick_fix
    r = m2.handle(ctx2, "ЮФУ")  # это ответ на вопрос, а не выбор поля для правки
    assert "student_name" not in ctx2.data  # «назад» вернул к первому вопросу и стёр прежний ответ
    assert r.kind in ("question", "error", "clarify")


def test_expired_data_is_purged_on_any_request_not_only_on_new_sessions(tmp_settings, tmp_path, monkeypatch):
    app = create_app(tmp_settings)
    admin = TestClient(app)
    admin_login(admin, app)
    admin.post("/api/admin/templates", files={"file": ("p.docx", build_sample_template(tmp_path / "t.docx").read_bytes(), "application/octet-stream")}, data={"template_id": "perevod"})
    c = TestClient(app)
    say(c, "Я хочу перевестись", "да", *ANSWERS, "Прикладная информатика", "1234 567890", "верно")
    c.post("/api/assistant/generate", json={"template_id": "perevod"})
    path = app.state.storage.db.query("SELECT path FROM generated_docs")[0]["path"]
    app.state.storage.db.execute("UPDATE generated_docs SET expires_at=?", (time.time() - 1,))
    real = time.time
    monkeypatch.setattr(time, "time", lambda: real() + 700)  # прошло больше 10 минут с прошлой чистки
    c.get("/api/assistant/state")  # обычный запрос существующей сессии, новых сессий нет
    assert not __import__("pathlib").Path(path).exists() and app.state.storage.db.query("SELECT * FROM generated_docs") == []


def test_fields_with_parts_cannot_be_edited_as_a_whole():
    from dekanat.dialog.docstate import apply_correction

    record = {"fields": [{"id": "issue", "title": "Номер и дата выдачи", "kind": "composite", "found": True, "value": "", "normalized": None, "confidence": 1.0, "method": "",
                          "parts": [{"id": "date", "title": "дата", "found": False, "value": "", "normalized": None, "confidence": 0.0}]}], "corrected": []}
    ok, error = apply_correction(record, None, "issue", "15.05.2024")
    assert not ok and "частями" in error.replace("части", "частями") and record["corrected"] == []


def test_concurrent_uploads_of_one_template_get_separate_version_files(tmp_path):
    from dekanat.templates import TemplateStore

    src = build_sample_template(tmp_path / "t.docx")
    store = TemplateStore(tmp_path / "templates")
    store.add(src, template_id="perevod")
    errors, infos = [], []

    def upload():
        try:
            infos.append(store.add(src, template_id="perevod", overwrite=True))
        except Exception as exc:  # noqa: BLE001
            errors.append(exc)

    threads = [threading.Thread(target=upload) for _ in range(6)]
    [t.start() for t in threads]
    [t.join() for t in threads]
    assert not errors and sorted(i.version for i in infos) == [2, 3, 4, 5, 6, 7]
    assert len({i.path for i in infos}) == 6 and all(i.path.exists() for i in infos)


# ============================================================================ на сайте нет технической информации
def _wait_job(c, job_id, timeout=60.0):
    end = time.time() + timeout
    while time.time() < end:
        snap = c.get(f"/api/jobs/{job_id}").json()
        if snap["status"] in ("done", "error", "cancelled"):
            return snap
        time.sleep(0.05)
    raise AssertionError("задача не завершилась")


def test_check_errors_are_shown_without_technical_details(tmp_settings, monkeypatch):
    """Посетитель видит понятный текст; название исключения и программы — только в журнале сервера."""
    from dekanat.ocr.engine import OcrUnavailable
    from dekanat.web.app import OCR_DOWN_TEXT

    app = create_app(tmp_settings)
    c = TestClient(app)
    for exc, expected in ((RuntimeError("KeyError в extract"), "Не удалось проверить документ"), (OcrUnavailable("Tesseract не найден: tesseract.exe"), OCR_DOWN_TEXT)):
        def boom(*_a, exc=exc, **_k):
            raise exc
        monkeypatch.setattr(app.state.svc, "analyze", boom)
        job = c.post("/api/jobs", data={"question": "Сколько дней рассматривается заявление?"}).json()
        snap = _wait_job(c, job["id"])
        assert snap["status"] == "error" and expected in snap["error"]
        for leaked in ("RuntimeError", "KeyError", "Tesseract", "tesseract.exe", "журнал"):
            assert leaked not in snap["error"]
        if isinstance(exc, OcrUnavailable):
            r = c.post("/api/analyze", data={"question": "?"})
            assert r.status_code == 503 and r.json()["detail"] == OCR_DOWN_TEXT


def test_reading_stage_labels_use_plain_words():
    """Этап «чтение текста» не показывает внутренние названия вариантов обработки (clahe, binary, page+deskew…)."""
    from dekanat.pipeline import _ocr_hint

    assert _ocr_hint("page+deskew · clahe") == "выравниваю освещение"
    assert _ocr_hint("deskew · binary") == "повышаю контраст"
    assert _ocr_hint("page · clean") == "нахожу лист"
    assert _ocr_hint("straight · clean") == ""
    assert _ocr_hint("проверяю поворот страницы") == "проверяю поворот страницы"
    assert _ocr_hint("rotate 90") == ""
    for geo in ("page", "page+deskew", "deskew", "straight"):
        for variant in ("clean", "binary", "deglare", "clahe"):
            hint = _ocr_hint(f"{geo} · {variant}")
            assert not any(word in hint for word in ("page", "deskew", "straight", "clean", "binary", "deglare", "clahe"))
