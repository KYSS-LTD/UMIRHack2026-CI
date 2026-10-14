"""DOCX-шаблоны: проверка файла, плейсхолдеры, заполнение абзацев/таблиц/колонтитулов, разбиение по runs, сохранность шаблона."""

import hashlib
import zipfile
from pathlib import Path

import pytest
from docx import Document

from dekanat.templates import (
    STANDARD_KEYS,
    TemplateError,
    TemplateExists,
    TemplateNotFound,
    TemplateStore,
    convert_to_template,
    fill_docx,
    find_placeholders,
    validate_docx,
)
from dekanat.templates.sample import build_enrollment_template, build_sample_template

DATA = {
    "student_name": "Иванов Иван Иванович", "passport_data": "серия 1234 № 567890", "current_university": "Южный федеральный университет", "current_course": "2",
    "current_specialty": "09.03.01 Информатика", "target_specialty": "09.03.04 Программная инженерия", "education_form": "очная", "debts": "отсутствуют",
    "application_date": "13.10.2026",
}


def sha(path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def make(path: Path, build) -> Path:
    doc = Document()
    build(doc)
    doc.save(str(path))
    return path


def all_text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    for s in doc.sections:
        parts += [p.text for p in s.header.paragraphs] + [p.text for p in s.footer.paragraphs]
    return "\n".join(parts)


@pytest.fixture()
def sample(tmp_path):
    return build_sample_template(tmp_path / "sample.docx")


# ============================================================================ образец «Заявление о зачислении»
def test_enrollment_sample_uses_only_the_questionnaire_keys_and_fills_completely(tmp_path):
    src = build_enrollment_template(tmp_path / "enroll.docx")
    assert validate_docx(src) is not None
    found = find_placeholders(src)
    assert {p.key for p in found} == set(STANDARD_KEYS) and all(p.known for p in found)
    res = fill_docx(src, DATA, tmp_path / "out")
    text = all_text(Document(str(res.path)))
    assert res.complete and "{{" not in text
    for value in ("Иванов Иван Иванович", "09.03.04 Программная инженерия", "очная", "Южный федеральный университет", "серия 1234 № 567890", "отсутствуют", "13.10.2026"):
        assert value in text
    assert "Прошу зачислить меня" in text


def test_enrollment_example_file_in_the_repo_is_the_builders_output(tmp_path):
    """Файл examples/templates/zayavlenie_o_zachislenii.docx — то, что записывает `template-sample --kind enrollment`: образец не расходится с кодом."""
    shipped = Path(__file__).resolve().parent.parent / "examples" / "templates" / "zayavlenie_o_zachislenii.docx"
    assert validate_docx(shipped) is not None
    built = build_enrollment_template(tmp_path / "built.docx")
    assert all_text(Document(str(shipped))) == all_text(Document(str(built)))
    assert {p.key for p in find_placeholders(shipped)} == set(STANDARD_KEYS)


def test_enrollment_sample_without_passport_leaves_a_blank_line(tmp_path):
    src = build_enrollment_template(tmp_path / "enroll.docx")
    res = fill_docx(src, {**DATA, "passport_data": None}, tmp_path / "out", blank_unfilled=True)
    text = all_text(Document(str(res.path)))
    assert "{{" not in text and "________" in text and res.unfilled == ["passport_data"]


# ============================================================================ плейсхолдеры
def test_all_standard_keys_are_found_in_the_sample(sample):
    found = {p.key: p for p in find_placeholders(sample)}
    assert set(STANDARD_KEYS) <= set(found) and all(p.known for p in found.values())
    assert found["student_name"].count == 1  # найден, хотя разбит на два run
    assert set(found["application_date"].locations) == {"table", "footer"} and found["passport_data"].locations == ["table"]


def test_placeholder_syntax_variants(tmp_path):
    def build(d):
        d.add_paragraph("{{a}} {{  b  }} {{ c_1 }} {{ Ключ }}")
        d.add_paragraph("{ one } {{ 1bad }} {{ with space }} {{}} {{ }}")
    found = find_placeholders(make(tmp_path / "t.docx", build))
    assert [p.key for p in found] == ["a", "b", "c_1", "Ключ"] and not any(p.known for p in found)


def test_custom_keys_are_discovered_and_filled(tmp_path):
    src = make(tmp_path / "t.docx", lambda d: d.add_paragraph("Группа: {{ group_name }}, ФИО: {{ student_name }}"))
    info = find_placeholders(src)
    assert [(p.key, p.known) for p in info] == [("group_name", False), ("student_name", True)]
    res = fill_docx(src, {"group_name": "ИВТ-21", "student_name": "Петров П. П."}, tmp_path / "out")
    assert Document(str(res.path)).paragraphs[0].text == "Группа: ИВТ-21, ФИО: Петров П. П." and res.complete


# ============================================================================ заполнение
def test_fill_paragraphs_tables_headers_footers(sample, tmp_path):
    res = fill_docx(sample, DATA, tmp_path / "out")
    doc = Document(str(res.path))
    text = all_text(doc)
    assert "{{" not in text and res.complete and res.unfilled == []
    assert "Иванов Иван Иванович" in doc.paragraphs[0].text            # плейсхолдер, разбитый на runs
    assert "из Южный федеральный университет" in text and "на 2 курсе" in text and "«09.03.04 Программная инженерия»" in text
    assert doc.tables[0].rows[0].cells[1].text == "серия 1234 № 567890"  # таблица
    assert doc.sections[0].footer.paragraphs[0].text == "Дата подачи: 13.10.2026"  # колонтитул
    assert "Академические задолженности: отсутствуют." in text
    assert set(res.filled) == set(STANDARD_KEYS)


def test_formatting_of_the_first_run_is_preserved(sample, tmp_path):
    doc = Document(str(fill_docx(sample, DATA, tmp_path / "out").path))
    runs = [r for r in doc.paragraphs[0].runs if r.text]
    name_run = next(r for r in runs if "Иванов" in r.text)
    assert name_run.bold is True  # шрифт run, где начинался плейсхолдер, сохранён
    assert doc.paragraphs[2].text.startswith("Прошу рассмотреть возможность моего перевода из")


def test_template_is_never_modified_and_outputs_are_unique(sample, tmp_path):
    before = sha(sample)
    a = fill_docx(sample, DATA, tmp_path / "out")
    b = fill_docx(sample, {**DATA, "student_name": "Петров Пётр Петрович"}, tmp_path / "out")
    assert sha(sample) == before and a.path != b.path and a.path.exists() and b.path.exists()
    assert "Иванов" in all_text(Document(str(a.path))) and "Петров" in all_text(Document(str(b.path)))
    assert a.path.parent == b.path.parent == tmp_path / "out" and a.path.parent != sample.parent


def test_missing_data_is_not_invented_and_is_reported(sample, tmp_path):
    data = {k: v for k, v in DATA.items() if k not in ("passport_data", "target_specialty")}
    res = fill_docx(sample, data, tmp_path / "out")
    assert not res.complete and set(res.unfilled) == {"passport_data", "target_specialty"}
    text = all_text(Document(str(res.path)))
    assert "{{ passport_data }}" in text and "{{ target_specialty }}" in text  # вместо выдуманного значения метка остаётся
    blank = fill_docx(sample, data, tmp_path / "out", blank_unfilled=True)
    assert "{{" not in all_text(Document(str(blank.path))) and set(blank.unfilled) == {"passport_data", "target_specialty"}


def test_empty_none_and_whitespace_values_count_as_missing(sample, tmp_path):
    res = fill_docx(sample, {**DATA, "passport_data": None, "debts": "   ", "target_specialty": ""}, tmp_path / "out")
    assert set(res.unfilled) == {"passport_data", "debts", "target_specialty"}


def test_unused_data_keys_are_listed(sample, tmp_path):
    res = fill_docx(sample, {**DATA, "nonexistent": "x"}, tmp_path / "out")
    assert res.unused == ["nonexistent"] and res.complete


def test_values_are_sanitised_and_not_expanded_again(tmp_path):
    src = make(tmp_path / "t.docx", lambda d: d.add_paragraph("A: {{ a }}; B: {{ b }}"))
    res = fill_docx(src, {"a": "строка1\nстрока2\t\x00конец", "b": "{{ a }}"}, tmp_path / "out")
    text = Document(str(res.path)).paragraphs[0].text
    assert text == "A: строка1 строка2 конец; B: {{ a }}"  # управляющие символы убраны, подставленное значение повторно не разворачивается


def test_same_placeholder_many_times_and_adjacent(tmp_path):
    src = make(tmp_path / "t.docx", lambda d: d.add_paragraph("{{ x }}{{ x }} и {{ y }}{{ x }}"))
    res = fill_docx(src, {"x": "1", "y": "2"}, tmp_path / "out")
    assert Document(str(res.path)).paragraphs[0].text == "11 и 21" and res.replaced == 4


def test_placeholder_split_over_three_runs_and_nested_table(tmp_path):
    def build(d):
        p = d.add_paragraph()
        for chunk in ("Студент {{", " student", "_name }}!"):
            p.add_run(chunk)
        outer = d.add_table(rows=1, cols=1)
        inner = outer.rows[0].cells[0].add_table(rows=1, cols=1)
        inner.rows[0].cells[0].text = "Курс: {{ current_course }}"
    res = fill_docx(make(tmp_path / "t.docx", build), DATA, tmp_path / "out")
    doc = Document(str(res.path))
    assert doc.paragraphs[0].text == "Студент Иванов Иван Иванович!"
    assert doc.tables[0].rows[0].cells[0].tables[0].rows[0].cells[0].text == "Курс: 2"


def test_filling_a_missing_template_is_a_clear_error(tmp_path):
    with pytest.raises(TemplateNotFound):
        fill_docx(tmp_path / "нет.docx", DATA, tmp_path / "out")


# ============================================================================ проверка файла и convert_to_template
def test_convert_returns_id_and_placeholder_info(sample, tmp_path):
    info = convert_to_template(sample, templates_dir=tmp_path / "templates", name="Заявление о переводе")
    assert info.template_id.startswith("zayavlenie-o-perevode-") and info.version == 1 and info.name == "Заявление о переводе"
    assert set(STANDARD_KEYS) <= set(info.keys) and info.custom_keys == []
    assert info.path == tmp_path / "templates" / info.template_id / "v1.docx" and info.path.exists() and sha(info.path) == sha(sample)
    d = info.to_dict()
    assert d["template_id"] == info.template_id and d["placeholders"][0]["key"]


@pytest.mark.parametrize("name, content, message", [
    ("нет.docx", None, "не найден"), ("a.txt", b"text", "должен быть файлом .docx"), ("a.doc", b"\xd0\xcf\x11\xe0", "должен быть файлом .docx"),
    ("a.docm", b"PK", "должен быть файлом .docx"), ("empty.docx", b"", "пустой"), ("broken.docx", b"this is not a zip", "повреждён"),
])
def test_invalid_files_are_rejected(tmp_path, name, content, message):
    path = tmp_path / name
    if content is not None:
        path.write_bytes(content)
    with pytest.raises(TemplateError, match=message):
        convert_to_template(path, templates_dir=tmp_path / "templates")
    assert not (tmp_path / "templates").exists() or not any((tmp_path / "templates").rglob("*.docx"))


def test_zip_without_word_document_is_rejected(tmp_path):
    p = tmp_path / "x.docx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("hello.txt", "hi")
    with pytest.raises(TemplateError, match="не документ Word"):
        validate_docx(p)


def _with_extra(src: Path, dst: Path, name: str, data: bytes) -> Path:
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr(name, data)
    return dst


def test_macros_and_external_templates_are_rejected(sample, tmp_path):
    with pytest.raises(TemplateError, match="макрос"):
        validate_docx(_with_extra(sample, tmp_path / "macro.docx", "word/vbaProject.bin", b"\x00evil"))
    rels = b'<?xml version="1.0"?><Relationships xmlns="x"><Relationship Id="r1" Type="http://x/attachedTemplate" Target="http://evil/t.dotm" TargetMode="External"/></Relationships>'
    with pytest.raises(TemplateError, match="внешний"):
        validate_docx(_with_extra(sample, tmp_path / "ext.docx", "word/_rels/settings.xml.rels", rels))


def test_zip_bomb_is_rejected(sample, tmp_path):
    with pytest.raises(TemplateError, match="архив-бомба|слишком большой"):
        validate_docx(_with_extra(sample, tmp_path / "bomb.docx", "word/media/big.bin", b"\x00" * (8 * 1024 * 1024)))


def test_size_limit(sample):
    with pytest.raises(TemplateError, match="больше"):
        validate_docx(sample, max_bytes=100)


def test_existing_template_is_not_overwritten_without_explicit_request(sample, tmp_path):
    root = tmp_path / "templates"
    first = convert_to_template(sample, templates_dir=root, template_id="perevod-2026")
    with pytest.raises(TemplateExists, match="перезапись"):
        convert_to_template(sample, templates_dir=root, template_id="perevod-2026")
    assert sha(first.path) == sha(sample)
    other = make(tmp_path / "other.docx", lambda d: d.add_paragraph("Другой {{ student_name }}"))
    replaced = convert_to_template(other, templates_dir=root, template_id="perevod-2026", overwrite=True)
    assert replaced.version == 2 and first.path.exists()  # прежняя версия остаётся в каталоге
    store = TemplateStore(root)
    assert store.path_of("perevod-2026") == replaced.path
    # без явного id каждый загруженный файл получает новый уникальный идентификатор и ничего не затирает
    ids = {convert_to_template(sample, templates_dir=root).template_id for _ in range(3)}
    assert len(ids) == 3 and "perevod-2026" not in ids


@pytest.mark.parametrize("bad_id", ["../evil", "A B", "x", "a/b", "..", "UPPER", "id with space", "a" * 80])
def test_unsafe_template_ids_are_rejected(sample, tmp_path, bad_id):
    with pytest.raises(TemplateError, match="Идентификатор"):
        convert_to_template(sample, templates_dir=tmp_path / "templates", template_id=bad_id)
    assert not list(tmp_path.rglob("evil*"))


def test_delete_removes_files_and_later_use_is_a_clear_error(sample, tmp_path):
    root = tmp_path / "templates"
    info = convert_to_template(sample, templates_dir=root)
    store = TemplateStore(root)
    assert store.exists(info.template_id) and store.delete(info.template_id) == 1
    assert not store.exists(info.template_id) and not (root / info.template_id).exists()
    with pytest.raises(TemplateNotFound, match="Шаблон не найден"):
        store.path_of(info.template_id)
    with pytest.raises(TemplateNotFound):
        store.delete(info.template_id)


def test_default_directory_comes_from_settings(sample, tmp_path, monkeypatch):
    from dekanat.config import Settings, load_settings

    base = load_settings()
    monkeypatch.setattr("dekanat.config.load_settings", lambda *a, **k: Settings(raw=base.raw, root=tmp_path))
    info = convert_to_template(sample)
    assert info.path.is_relative_to(tmp_path / "templates")  # внутренний каталог /templates/ из ТЗ
